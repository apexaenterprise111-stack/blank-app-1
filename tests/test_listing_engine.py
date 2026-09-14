from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from listing_engine import generate_customer_workbook, inspect_workbook


class ListingEngineTests(unittest.TestCase):
    def master_bytes(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Catalog"
        sheet.append(["Current product master"])
        sheet.append(
            [
                "Group ID",
                "SKU",
                "Parentage",
                "Product Name",
                "Color",
                "Pack",
                "Size",
                "MRP",
                "Main Image",
                "Style Code",
                "Title",
                "Description",
                "Search Keywords",
                "Bullet Point 1",
                "Bullet Point 2",
                "Bullet Point 3",
                "Bullet Point 4",
                "Bullet Point 5",
            ]
        )
        rows = [
            ["Group 02", "P2", "Parent", "COTTON CAMISOLE", "Blue", 2, "S", 99, "https://example/2.jpg", "ST2", "", "", "", "", "", "", "", ""],
            ["Group 02", "C2", "Child", "COTTON CAMISOLE", "Blue", 2, "M", 99, "https://example/2.jpg", "ST2", "", "", "", "", "", "", "", ""],
            ["Group 01", "P1", "Parent", "COTTON CAMISOLE", "Red", 1, "S", 89, "https://example/1.jpg", "ST1", "", "", "", "", "", "", "", ""],
            ["Group 01", "C1", "Child", "COTTON CAMISOLE", "Red", 1, "M", 89, "https://example/1.jpg", "ST1", "", "", "", "", "", "", "", ""],
        ]
        for row in rows:
            sheet.append(row)
        validation = DataValidation(type="list", formula1='"S,M"')
        validation.add("G3:G6")
        sheet.add_data_validation(validation)
        notes = workbook.create_sheet("Notes")
        notes["A1"] = "Must remain"
        buffer = BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def test_inspection_uses_current_workbook(self):
        profile = inspect_workbook(self.master_bytes(), "current.xlsx")
        self.assertIsNone(profile["error"])
        self.assertEqual(profile["recommended_sheet"], "Catalog")
        catalog = next(sheet for sheet in profile["sheets"] if sheet["name"] == "Catalog")
        self.assertIn("Product Name", [column["header"] for column in catalog["columns"]])
        self.assertIn("Title", catalog["content_fields"])
        self.assertEqual(catalog["product_identities"], ["COTTON CAMISOLE"])

    def test_ten_copies_preserve_locked_values_and_group_order(self):
        source = self.master_bytes()
        results = [
            generate_customer_workbook(source, "current.xlsx", "Amazon", customer, "Catalog")
            for customer in range(1, 11)
        ]
        self.assertTrue(all(result.success for result in results), [result.errors for result in results])
        self.assertEqual(len({load_workbook(BytesIO(result.data))["Catalog"]["K3"].value for result in results}), 10)

        output = load_workbook(BytesIO(results[0].data), data_only=False)
        sheet = output["Catalog"]
        self.assertEqual([sheet[f"A{row}"].value for row in range(3, 7)], ["Group 01", "Group 01", "Group 02", "Group 02"])
        self.assertEqual([sheet[f"B{row}"].value for row in range(3, 7)], ["P1", "C1", "P2", "C2"])
        self.assertEqual([sheet[f"H{row}"].value for row in range(3, 7)], [89, 89, 99, 99])
        self.assertEqual([sheet[f"I{row}"].value for row in range(3, 7)], ["https://example/1.jpg", "https://example/1.jpg", "https://example/2.jpg", "https://example/2.jpg"])
        self.assertEqual([sheet[f"J{row}"].value for row in range(3, 7)], ["ST1", "ST1", "ST2", "ST2"])
        self.assertEqual(len(sheet.data_validations.dataValidation), 1)
        self.assertEqual(output.sheetnames, ["Catalog", "Notes"])

    def test_flipkart_description_does_not_include_prohibited_brand(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Flipkart"
        sheet.append(["Group ID", "SKU", "Product Name", "Description", "Title"])
        sheet.append(["1", "SKU-1", "KSHTABHANJAN Camisole", "", ""])
        buffer = BytesIO()
        workbook.save(buffer)
        result = generate_customer_workbook(buffer.getvalue(), "flip.xlsx", "Flipkart", 1, "Flipkart")
        self.assertTrue(result.success, result.errors)
        output = load_workbook(BytesIO(result.data))
        self.assertNotIn("KSHTABHANJAN", output["Flipkart"]["D2"].value.upper())


if __name__ == "__main__":
    unittest.main()
